import openpyxl

wb = openpyxl.load_workbook("website_analysis_results.xlsx")

# 1. Assert sheet names
expected_sheets = ["Qualified", "Rejected", "Manual Review"]
assert wb.sheetnames == expected_sheets, f"Sheets mismatch: {wb.sheetnames} vs {expected_sheets}"
print("Sheets check passed:", wb.sheetnames)

# 2. Check headers
ws_qual = wb["Qualified"]
qual_headers = [c.value for c in ws_qual[1]]
expected_qual = [
    "No", "Website Name", "Domain", "Website URL", "Login URL", "Registration URL",
    "Google Auth Entry URL", "Google Auth Status", "Member Area URL", "Profile URL",
    "Edit Profile URL", "Checked Date"
]
assert qual_headers == expected_qual, f"Qual headers mismatch: {qual_headers}"

ws_rej = wb["Rejected"]
rej_headers = [c.value for c in ws_rej[1]]
expected_rej = ["No", "Website Name", "Domain", "Website URL", "Reason"]
assert rej_headers == expected_rej, f"Rej headers mismatch: {rej_headers}"

ws_man = wb["Manual Review"]
man_headers = [c.value for c in ws_man[1]]
expected_man = ["No", "Website Name", "Domain", "Website URL", "Reason"]
assert man_headers == expected_man, f"Man headers mismatch: {man_headers}"
print("Headers check passed!")

# 3. Check row counts
q_rows = ws_qual.max_row - 1
r_rows = ws_rej.max_row - 1
m_rows = ws_man.max_row - 1
total_rows = q_rows + r_rows + m_rows
print(f"Row counts: Qualified={q_rows}, Rejected={r_rows}, Manual Review={m_rows}, Total={total_rows}")
assert total_rows == 122, f"Total rows is {total_rows}, expected 122!"

# 4. Check hyperlinks
for ws_name, url_cols in [("Qualified", [4, 5, 6, 7, 9, 10, 11]), ("Rejected", [4]), ("Manual Review", [4])]:
    ws = wb[ws_name]
    for r in range(2, ws.max_row + 1):
        for c in url_cols:
            cell = ws.cell(row=r, column=c)
            val = str(cell.value or "")
            if val.startswith("http://") or val.startswith("https://"):
                assert cell.hyperlink is not None, f"Missing hyperlink in {ws_name} at row {r}, col {c}: {val}"
                assert cell.hyperlink.target == val, f"Hyperlink mismatch: {cell.hyperlink.target} vs {val}"

print("All hyperlink checks passed successfully!")
print("VERIFICATION COMPLETED WITH 100% SUCCESS!")
