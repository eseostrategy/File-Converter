import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open("analysis_results.json") as f:
    data = json.load(f)

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# 1. Qualified sheet
# Columns: No, Website Name, Domain, Website URL, Login URL, Registration URL, Google Auth Entry URL, Google Auth Status (Verified + short evidence), Member Area URL, Profile URL, Edit Profile URL, Checked Date
ws_qual = wb.create_sheet(title="Qualified")
qual_headers = [
    "No",
    "Website Name",
    "Domain",
    "Website URL",
    "Login URL",
    "Registration URL",
    "Google Auth Entry URL",
    "Google Auth Status",
    "Member Area URL",
    "Profile URL",
    "Edit Profile URL",
    "Checked Date"
]
ws_qual.append(qual_headers)

# 2. Rejected sheet
# Columns: No, Website Name, Domain, Website URL, Reason
ws_rej = wb.create_sheet(title="Rejected")
rej_headers = [
    "No",
    "Website Name",
    "Domain",
    "Website URL",
    "Reason"
]
ws_rej.append(rej_headers)

# 3. Manual Review sheet
# Columns: No, Website Name, Domain, Website URL, Reason
ws_man = wb.create_sheet(title="Manual Review")
man_headers = [
    "No",
    "Website Name",
    "Domain",
    "Website URL",
    "Reason"
]
ws_man.append(man_headers)

link_font = Font(color="0000FF", underline="single", name="Calibri", size=11)
normal_font = Font(name="Calibri", size=11)
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

header_fills = {
    ws_qual: PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid"),
    ws_rej: PatternFill(start_color="C62828", end_color="C62828", fill_type="solid"),
    ws_man: PatternFill(start_color="E65100", end_color="E65100", fill_type="solid")
}

def set_hyperlink(cell, url_str):
    if not url_str:
        cell.value = "N/A"
        cell.font = normal_font
        return
    url_str = str(url_str).strip()
    if url_str.startswith("http://") or url_str.startswith("https://"):
        cell.value = url_str
        cell.hyperlink = url_str
        cell.font = link_font
    else:
        cell.value = url_str
        cell.font = normal_font

qual_idx = 1
rej_idx = 1
man_idx = 1

for item in data:
    status = item.get("status")
    if status == "QUALIFIED":
        row_vals = [
            qual_idx,
            item.get("website_name", ""),
            item.get("domain", ""),
            item.get("website_url", ""),
            item.get("login_url", ""),
            item.get("registration_url", ""),
            item.get("google_entry_url", ""),
            item.get("google_auth_status", ""),
            item.get("member_area_url", "Not visible without login"),
            item.get("profile_url", "Not visible without login"),
            item.get("edit_profile_url", "Not visible without login"),
            item.get("checked_date", "2026-11-01")
        ]
        ws_qual.append(row_vals)
        cur_row = ws_qual.max_row
        # apply hyperlinks to URL columns: 4, 5, 6, 7, 9, 10, 11
        for col_idx in [4, 5, 6, 7, 9, 10, 11]:
            cell = ws_qual.cell(row=cur_row, column=col_idx)
            set_hyperlink(cell, cell.value)
        qual_idx += 1

    elif status == "REJECTED":
        row_vals = [
            rej_idx,
            item.get("website_name", ""),
            item.get("domain", ""),
            item.get("website_url", ""),
            item.get("reason", "Page loaded successfully and has no Google option.")
        ]
        ws_rej.append(row_vals)
        cur_row = ws_rej.max_row
        # apply hyperlink to Website URL (col 4)
        cell = ws_rej.cell(row=cur_row, column=4)
        set_hyperlink(cell, cell.value)
        rej_idx += 1

    else: # MANUAL_REVIEW
        row_vals = [
            man_idx,
            item.get("website_name", ""),
            item.get("domain", ""),
            item.get("website_url", ""),
            item.get("reason", "Manual review required.")
        ]
        ws_man.append(row_vals)
        cur_row = ws_man.max_row
        # apply hyperlink to Website URL (col 4)
        cell = ws_man.cell(row=cur_row, column=4)
        set_hyperlink(cell, cell.value)
        man_idx += 1

# Style headers & auto-fit columns
thin_border = Border(
    left=Side(style='thin', color='D0D0D0'),
    right=Side(style='thin', color='D0D0D0'),
    top=Side(style='thin', color='D0D0D0'),
    bottom=Side(style='thin', color='D0D0D0')
)

for ws in [ws_qual, ws_rej, ws_man]:
    # Header styling
    for col_num in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fills[ws]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Row styling & column width
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if not cell.hyperlink:
                cell.font = normal_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if cell.hyperlink:
                # display length might be long, cap at 45 for width
                max_len = max(max_len, min(len(val_str), 45))
            else:
                max_len = max(max_len, min(len(val_str), 60))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save("website_analysis_results.xlsx")
print("Saved website_analysis_results.xlsx successfully!")
